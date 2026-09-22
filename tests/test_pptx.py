from pathlib import Path
from zipfile import ZipFile

from image2svg.core.scene import Scene, SceneElement
from image2svg.export.pptx import export_pptx


def test_pptx_exports_primitives_and_connector_as_native_objects(tmp_path: Path) -> None:
    from pptx import Presentation

    scene = Scene(
        width=300,
        height=200,
        background="#FFFFFF",
        elements=[
            SceneElement(
                id="rect",
                type="rect",
                bbox=(10, 10, 80, 40),
                style={"fill": "#DDEEFF", "stroke": "#112233"},
            ),
            SceneElement(
                id="ellipse",
                type="ellipse",
                bbox=(110, 10, 50, 40),
                style={"fill": "#44AA88"},
            ),
            SceneElement(
                id="triangle",
                type="polygon",
                bbox=(180, 10, 60, 50),
                points=[(180, 60), (210, 10), (240, 60)],
                style={"fill": "#FFAA00"},
            ),
            SceneElement(
                id="arrow",
                type="group",
                bbox=(90, 30, 20, 0),
                style={
                    "arrow_start": [90.0, 30.0],
                    "arrow_end": [110.0, 30.0],
                    "stroke": "#123456",
                    "stroke_width": 2.0,
                    "arrowhead": True,
                },
            ),
            SceneElement(
                id="text",
                type="text",
                bbox=(20, 100, 260, 40),
                text="可编辑文字",
                style={
                    "font-family": "Microsoft YaHei, Arial",
                    "font-size": 24,
                    "font-weight": "bold",
                    "fill": "#123456",
                    "align": "center",
                    "fit-to-box": True,
                },
            ),
        ],
    )
    output = export_pptx(scene, tmp_path / "native.pptx")

    presentation = Presentation(output)
    xml = presentation.slides[0].element.xml
    assert xml.count("<p:cxnSp>") == 1
    assert "<a:tailEnd" in xml
    assert "<p:pic>" not in xml
    text_shape = next(
        shape for shape in presentation.slides[0].shapes if shape.has_text_frame and shape.text
    )
    run = text_shape.text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Microsoft YaHei"
    assert run.font.bold is True
    assert "<a:normAutofit" in text_shape.text_frame._txBody.xml


def test_pptx_rasterizes_vector_groups_at_double_resolution(
    tmp_path: Path, monkeypatch
) -> None:
    from PIL import Image

    def fake_render(svg: str, output_path: Path, *, scale: float = 1.0) -> Path:
        assert "width=\"40\"" in svg
        Image.new("RGBA", (round(40 * scale), round(30 * scale))).save(output_path)
        return output_path

    monkeypatch.setattr("image2svg.svg.render.render_svg", fake_render)

    scene = Scene(
        width=100,
        height=100,
        elements=[
            SceneElement(
                id="icon",
                type="group",
                bbox=(10, 20, 40, 30),
                raw_svg='<rect x="0" y="0" width="40" height="30" fill="#3366CC"/>',
            )
        ],
    )

    output = export_pptx(scene, tmp_path / "vector-group.pptx")

    with ZipFile(output) as archive:
        media_name = next(name for name in archive.namelist() if name.startswith("ppt/media/"))
        with archive.open(media_name) as stream, Image.open(stream) as image:
            assert image.size == (80, 60)
